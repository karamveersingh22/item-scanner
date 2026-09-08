import sqlite3
import openpyxl
import time
import os

EXCEL_FILE = "itemmast.xlsx"
DATABASE_FILE = "item_scanner.db"

start = time.time()

print("Opening Excel file...")

# Read Excel
workbook = openpyxl.load_workbook(
    EXCEL_FILE,
    read_only=True,
    data_only=True
)

sheet = workbook.active

print(f"Worksheet: {sheet.title}")
print("Reading headers...")

# Read header row
headers = next(sheet.iter_rows(values_only=True))

headers = [
    str(header).strip().upper() if header is not None else ""
    for header in headers
]

print("Columns found:")
print(headers)

required_columns = [
    "I_CODE",
    "ITEM_NAME",
    "DESCRIBE",
    "QUANTITY",
    "RATE",
    "DISC_PER",
]

# Find column positions
column_indexes = {}

for column in required_columns:
    if column not in headers:
        raise Exception(
            f"Required column '{column}' was not found."
        )

    column_indexes[column] = headers.index(column)

# DISC_B is optional/flexible
if "DISC_B" in headers:
    column_indexes["DISC_B"] = headers.index("DISC_B")

print("All required columns found.")

# Remove old database
if os.path.exists(DATABASE_FILE):
    os.remove(DATABASE_FILE)

print("Creating SQLite database...")

connection = sqlite3.connect(DATABASE_FILE)

cursor = connection.cursor()

cursor.execute("""
    CREATE TABLE items (
        I_CODE TEXT PRIMARY KEY,
        ITEM_NAME TEXT,
        DESCRIBE TEXT,
        QUANTITY TEXT,
        RATE TEXT,
        DISC_PER TEXT,
        DISC_B TEXT
    )
""")

# Prepare insert
insert_sql = """
    INSERT OR REPLACE INTO items
    (
        I_CODE,
        ITEM_NAME,
        DESCRIBE,
        QUANTITY,
        RATE,
        DISC_PER,
        DISC_B
    )
    VALUES (?, ?, ?, ?, ?, ?, ?)
"""

batch = []
count = 0

print("Converting Excel → SQLite...")

for row in sheet.iter_rows(
    min_row=2,
    values_only=True
):

    def value(column):
        if column not in column_indexes:
            return ""
        index = column_indexes[column]

        if index >= len(row):
            return ""

        cell = row[index]

        if cell is None:
            return ""

        return str(cell).strip()

    i_code = value("I_CODE")

    if not i_code:
        continue

    batch.append(
        (
            i_code,
            value("ITEM_NAME"),
            value("DESCRIBE"),
            value("QUANTITY"),
            value("RATE"),
            value("DISC_PER"),
            value("DISC_B"),
        )
    )

    count += 1

    # Insert in batches
    if len(batch) >= 1000:
        cursor.executemany(insert_sql, batch)
        batch.clear()

        print(f"Imported {count} items...")

# Insert remaining rows
if batch:
    cursor.executemany(insert_sql, batch)

connection.commit()

# Verify
cursor.execute("SELECT COUNT(*) FROM items")

database_count = cursor.fetchone()[0]

connection.close()

workbook.close()

elapsed = time.time() - start

print()
print("==============================")
print("CONVERSION COMPLETE")
print("==============================")
print(f"Excel records : {count}")
print(f"Database rows : {database_count}")
print(f"Time          : {elapsed:.2f} seconds")
print(f"Database      : {DATABASE_FILE}")
print("==============================")